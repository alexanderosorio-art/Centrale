import unittest
from io import BytesIO
from datetime import datetime
import openpyxl
from test_providers import FECHA
from processing import procesar_archivos
from providers.nexsys import leer, es_catalogo


def workbook(sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    stream.name = 'Nexsys Hardware.xlsx'
    return stream


class NexsysCatalogTests(unittest.TestCase):
    def test_currency_filter_and_sheet_allowlist(self):
        generic = [['SKU', 'Descripcion', 'PRECIO USD', 'STOCK'], ['UNWANTED', 'Producto', 10, 5]]
        stream = workbook({
            ' Retail ': [
                ['Código de artículo', 'Descripción Resumida', 'Precio Canal USD-CLP', 'Moneda', 'Stock'],
                ['A', 'Producto', 100, 'USD', 2], ['B', 'Producto', 100000, 'CLP', 2],
                ['C', 'Producto', 100, None, 2], ['D', 'Producto', 'CLP 99.800', 'USD', 2]],
            'Impresión Láser - Tinta': [
                ['Part Number', 'Descripcion', 'Precio USD', 'Stock'],
                ['E', 'Producto', 'CLP 122.500', 3], ['F', 'Producto', 150, 3]],
            'EPSON': generic, 'OFERTAS S4 SEPTIEMBRE': generic, 'Otra': generic,
        })
        summary, result = procesar_archivos([stream], FECHA, 'Nexsys', 'CLP')
        self.assertEqual(list(result.provider_code), ['A', 'F'])
        self.assertEqual(set(result.currency), {'USD'})
        self.assertIn('EPSON: excluida completa', summary.iloc[0]['Detalle'])
        self.assertIn('1 moneda sin identificar', summary.iloc[0]['Detalle'])

    def test_poly_variable_promotion_title(self):
        for title in [datetime(2026, 10, 31), 'Promoción noviembre', 'Nuevo precio']:
            with self.subTest(title=title):
                stream = workbook({' HP Poly ': [
                    ['Código de artículo', 'Nombre del artículo', 'Descripción', 'Inventario físico', title, 'Precio de venta US$'],
                    ['A', 'Nombre', 'Detalle', 10, 27.74, 38],
                    ['B', 'Nombre', 'Detalle', 1, None, 43],
                    ['C', 'Nombre', 'Detalle', 1, 0, 43],
                    ['D', 'Nombre', 'Detalle', 1, 'CLP 99', 43],
                    ['E', 'Nombre', 'Detalle', 'En transito 30 dias', 20, 43]]})
                result, _, _ = leer(stream, FECHA, 'USD')
                self.assertEqual(list(result.currency_unaware_cost_neto), [27.74, 43, 43, 43])

    def test_lenovo_promotion_quantity(self):
        stream = workbook({' LENOVO NOTEBOOK_WS ': [
            ['SKU', 'DESCRIPCION', 'PRECIO Lista NETO USD', 'Stock', 'Precio Promo Neto', 'Unidades Promo'],
            ['A', 'Producto', 780, '100+', 720, 50],
            ['B', 'Producto', 1650, '100+', 1500, 0],
            ['C', 'Producto', 977, '100+', 920, 162],
            ['D', 'Producto', 500, 10, 600, 2],
            ['E', 'Producto', 500, 0, 400, 2],
            ['F', 'Producto', 500, 10, 400, None]]})
        result, _, _ = leer(stream, FECHA, 'USD')
        self.assertEqual(list(zip(result.currency_unaware_cost_neto, result.quantity)),
                         [(720, 50), (1650, 100), (920, 100), (500, 10), (500, 10)])

    def test_description_not_warranty_and_no_connect_price(self):
        stream = workbook({'HP Notebook': [
            ['P/N', 'Precio Oferta USD', 'Precio Normal USD', 'QTY', 'PRECIO HP CONNECT', 'Descripcion ESS', 'Descripcion'],
            ['A', 419, 429, '(+) 200', 100, 'Garantia adicional', 'Equipo'],
            ['SKU', 'P/N', 'Descripcion ESS', 'Costo US$'],
            ['WARRANTY', 'OTHER-PN', 'Garantia adicional', 100]]})
        result, _, _ = leer(stream, FECHA, 'USD')
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0]['name'], 'Equipo')
        self.assertEqual(result.iloc[0].quantity, 200)
        self.assertEqual(result.iloc[0].currency_unaware_cost_neto, 419)
        self.assertEqual(result.iloc[0].brand, '')

    def test_jabra_certified_and_cell_currency(self):
        stream = workbook({'JABRA ': [
            ['SKU', 'Descripción', 'Precio Reseller Premium (certificados)', 'Stock', 'Tránsito'],
            ['A', 'Auricular', 24.9, 205, 100], ['B', 'Auricular', 24900, 10, 0]]})
        book = openpyxl.load_workbook(stream)
        book.active['C3'].number_format = '"CLP" #,##0'
        changed = BytesIO(); book.save(changed); book.close()
        changed.name = stream.name; changed.seek(0)
        result, _, _ = leer(changed, FECHA, 'USD')
        self.assertEqual(list(result.provider_code), ['A'])
        self.assertEqual(result.iloc[0].quantity, 205)

    def test_selection_and_asus_compatibility(self):
        self.assertTrue(es_catalogo([' hp poly ']))
        self.assertFalse(es_catalogo(['Notebook', 'Dekstop - AIO', 'Garantias']))
        stream = workbook({'HP POS': [['Código de artículo', 'Nombre del artículo', 'Precio Normal', 'Stock'], ['A', 'Equipo', 'USD 82', 2]]})
        with self.assertRaises(ValueError):
            leer(stream, FECHA, 'USD', hojas_seleccionadas=[])


if __name__ == '__main__':
    unittest.main()
