"""Pruebas con archivos sintéticos; no incluyen listas comerciales privadas."""
import sys
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from processing import procesar_archivos, moneda_predeterminada, detectar_proveedor_lote
from providers import REGISTRO
from common import consolidar
from exports import generar_excel, generar_html_copia_excel, generar_tsv

FECHA = date(2026, 9, 30)
PLANTILLAS = {
    'Intcomex': ['SKU', 'venta neto usd', 'stock actual'],
    'Kepler': ['CODIGO', 'VALOR', 'STOCK', 'DESCRIPCION'],
    'Tecnoglobal': ['CODIGO TG', 'PRECIO', 'STOCK', 'DESCRIPCION'],
    'Ingram': ['SKU INGRAM', 'COSTO', 'STOCK'],
    'Coimco': ['CODIGO', 'PRECIO', 'STOCK', 'DESCRIPCION'],
    'Fujicorp': ['CODIGO', 'PRECIO', 'STOCK', 'DESCRIPCION'],
    'Nexsys': ['SKU', 'PRECIO', 'STOCK', 'DESCRIPCION'],
    'SolutionBox': ['PN', 'PRECIO LISTA', 'STOCK', 'DESCRIPCION'],
    'Demco Ltda.': ['Código Interno', 'Precio Neto', 'Stock', 'Nombre'],
}


def excel(headers, rows, name='lista.xlsx', second_sheet=False):
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        pd.DataFrame(rows, columns=headers).to_excel(writer, index=False, sheet_name='MATERIAL')
        if second_sheet:
            pd.DataFrame(rows, columns=headers).to_excel(writer, index=False, sheet_name='OTRA')
    stream.name = name
    stream.seek(0)
    return stream


class ProviderTests(unittest.TestCase):
    def test_all_configured_providers_and_currency_override(self):
        for provider, headers in PLANTILLAS.items():
            for currency in (None, 'USD', 'CLP'):
                with self.subTest(provider=provider, currency=currency):
                    rows = [[code, price, stock] + (['Producto'] if len(headers) == 4 else [])
                            for code, price, stock in [('001', 12.75, 3), ('002', 15, 0), ('003', 0, 5)]]
                    summary, result = procesar_archivos([excel(headers, rows)], FECHA, provider, currency)
                    self.assertEqual(summary.iloc[0]['Estado'], 'Procesado', summary.to_dict())
                    self.assertEqual(len(result), 1)
                    self.assertEqual(result.iloc[0].currency_unaware_cost_neto, 12.75)
                    self.assertEqual(result.iloc[0].quantity, 3)
                    self.assertEqual(result.iloc[0].currency, currency or moneda_predeterminada(provider))
                    self.assertEqual(result.iloc[0].brand, '')
                    self.assertEqual(result.iloc[0].expiry_date, '30-09-2026  12:00')

    def test_no_provider_imports_another_provider(self):
        import ast
        root = Path(__file__).resolve().parents[1] / 'providers'
        for path in root.glob('*.py'):
            if path.name == '__init__.py':
                continue
            for node in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
                if isinstance(node, ast.ImportFrom):
                    self.assertFalse(node.level or (node.module or '').startswith('providers'), path.name)

    def test_defaults_and_detection(self):
        for provider, module in REGISTRO.items():
            self.assertEqual(moneda_predeterminada(provider), 'CLP' if provider in
                             ['Fujicorp', 'Coimco', 'Demco Ltda.', 'Facciatech', 'Gerona'] else 'USD')
            if module.PATRON:
                self.assertEqual(detectar_proveedor_lote([provider + ' 2026.xlsx']), (provider, 'detectado'))
        self.assertEqual(detectar_proveedor_lote(['Ingram.xlsx', 'Intcomex.xlsx'])[1], 'conflicto')

    def test_selected_sheets(self):
        for provider in ['Ingram', 'SolutionBox']:
            headers = PLANTILLAS[provider]
            row = ['001', 12, 3] + (['Producto'] if len(headers) == 4 else [])
            for selected, expected in [(['MATERIAL'], 1), (['MATERIAL', 'OTRA'], 2), ([], 0)]:
                summary, result = procesar_archivos([excel(headers, [row], second_sheet=True)],
                    FECHA, provider, hojas_por_archivo={0: selected})
                self.assertEqual(len(result), expected)
                self.assertEqual(summary.iloc[0]['Estado'], 'Procesado' if selected else 'Error')

    def test_fujicorp_offer_isolation(self):
        for provider, expected in [('Fujicorp', 8), ('Coimco', 12)]:
            headers = [*PLANTILLAS[provider], 'OFERTA', 'MARCA', 'PART NUMBER']
            _, result = procesar_archivos([excel(headers, [['A', 12, '10+', 'Producto', 8, 'Marca', 'PN-A']])], FECHA, provider)
            self.assertEqual(result.iloc[0].currency_unaware_cost_neto, expected)
            self.assertEqual(result.iloc[0].brand, 'Marca')
            self.assertEqual(result.iloc[0].mpn, 'PN-A')

    def test_demco_stock_and_price(self):
        headers = [*PLANTILLAS['Demco Ltda.'], 'PVP', 'Número de Parte', 'Marca']
        rows = [['A', 12, 5, 'Producto', 100, 'PN-A', 'Marca'], ['B', 12, '30 DIAS', 'Producto', 100, 'PN-B', 'Marca']]
        _, result = procesar_archivos([excel(headers, rows)], FECHA, 'Demco Ltda.')
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0].currency_unaware_cost_neto, 12)
        self.assertEqual(result.iloc[0].mpn, 'PN-A')

    def test_nexsys_pn_equivalence_already_authorized(self):
        for permission in [False, True]:
            summary, result = procesar_archivos([excel(['PART NUMBER', 'PRECIO', 'STOCK', 'DESCRIPCION'],
                [['PN-A', 12, 2, 'Producto']])], FECHA, 'Nexsys', permitir_pn=permission)
            self.assertEqual(len(result), 1)
            self.assertEqual(result.iloc[0].provider_code, result.iloc[0].mpn)
            self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')

    def test_kepler_preventa_and_explicit_fields(self):
        headers = [*PLANTILLAS['Kepler'], 'PN', 'MARCA']
        summary, result = procesar_archivos([excel(headers, [['A-\nB', 12, 3, 'Una\n descripción', 'PN-A', 'Marca']],
            name='PREVENTA TT.xlsx')], FECHA, 'Kepler')
        self.assertEqual(result.iloc[0].provider_code, 'A-B')
        self.assertEqual(result.iloc[0]['name'], 'Una descripción')
        self.assertEqual(result.iloc[0].mpn, 'PN-A')
        self.assertIn('Preventa', summary.iloc[0]['Detalle'])

    def test_errors_do_not_block_other_files(self):
        bad = BytesIO(b'not excel'); bad.name = 'bad.xlsx'
        good = excel(PLANTILLAS['Intcomex'], [['A', 1, 2]])
        summary, result = procesar_archivos([bad, good], FECHA)
        self.assertEqual(list(summary.Estado), ['Error', 'Procesado'])
        self.assertEqual(len(result), 1)

    def test_intcomex_hikvision(self):
        headers = ['Imagen', 'SKU', 'PART #', 'Descripción', 'DPV', 'Unit Price$']
        rows = [[None, 'ES213HIK97', 'PN-A', 'Cámara', '500+', 10.744],
                [None, 'B', 'PN-B', 'Otra', '20+', 22.969],
                [None, None, None, None, 'DPV', 'Unit Price$'],
                [None, 'C', 'PN-C', 'Sin stock', 0, 10],
                [None, 'D', 'PN-D', 'Comentario', '30 DIAS', 10]]
        summary, result = procesar_archivos([excel(headers, rows)], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(list(result.quantity), [500, 20])
        self.assertEqual(list(result.currency_unaware_cost_neto), [10.744, 22.969])
        self.assertEqual(result.iloc[0].mpn, 'PN-A')
        self.assertEqual(result.iloc[0]['name'], 'Cámara')
        self.assertTrue((result.brand == '').all())

    def test_exports_and_conflicts(self):
        _, result = procesar_archivos([excel(PLANTILLAS['Intcomex'], [['A', 12.75, 3], ['A', 12.75, 3], ['B', 5, 1], ['B', 6, 1]])], FECHA)
        clean, conflicts, duplicates = consolidar(result)
        self.assertEqual((len(clean), len(conflicts), duplicates), (1, 2, 1))
        data = clean.drop(columns='Archivo de origen')
        text = generar_tsv(data)
        self.assertNotIn('provider_code', text)
        self.assertIn('12,75', text)
        self.assertNotIn('3.0', text)
        html = generar_html_copia_excel(data)
        self.assertIn('x:str="&#x27;30-09-2026  12:00"', html)
        self.assertNotIn("''30", html)
        book = openpyxl.load_workbook(generar_excel(data))
        self.assertEqual(book.active['D2'].value, '30-09-2026  12:00')
        self.assertTrue(book.active['D2'].quotePrefix)
        book.close()

    def test_unconfigured_providers_are_not_guessed(self):
        for provider in ['Otro']:
            summary, result = procesar_archivos([excel(PLANTILLAS['Intcomex'], [['A', 1, 2]])], FECHA, provider)
            self.assertEqual(summary.iloc[0]['Estado'], 'Error')
            self.assertTrue(result.empty)

    def test_gerona_empty_stock(self):
        headers = ['CodigoSKU', 'Modelo', 'Descripcion', 'Marca', 'Precio Lista (sin IVA)', 'Precio Sugerido']
        rows = [[112275, 'CD289-90549', 'Cargador', 'UGREEN', 43692.0168067227, 79990],
                [112276, 'PN-B', 'OPEN BOX cargador', 'UGREEN', 100, 200],
                [112277, 'PN-C', 'Sin costo', None, None, 200]]
        summary, df = procesar_archivos([excel(headers, rows)], FECHA, 'Gerona')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0].provider_code, '112275')
        self.assertAlmostEqual(df.iloc[0].currency_unaware_cost_neto, 43692.0168067227)
        self.assertEqual(df.iloc[0].quantity, '')
        self.assertEqual(df.iloc[0].mpn, 'CD289-90549')
        self.assertEqual(df.iloc[0].currency, 'CLP')
        copia = generar_tsv(df.drop(columns='Archivo de origen')).splitlines()[0].split('\t')
        self.assertEqual(copia[4], '')
        libro = openpyxl.load_workbook(generar_excel(df.drop(columns='Archivo de origen')))
        self.assertIsNone(libro.active['E2'].value)

    def test_supplier_short_names(self):
        for nombre, esperado in [('LP_DELL_IX.xlsx', 'Intcomex'), ('LP_TG.xlsx', 'Tecnoglobal'), ('LP_IM.xlsx', 'Ingram')]:
            self.assertEqual(detectar_proveedor_lote([nombre]), (esperado, 'detectado'))
        self.assertEqual(detectar_proveedor_lote(['MAXIM.xlsx'])[1], 'sin_coincidencia')

    def test_intcomex_dell_current_stock_only(self):
        headers = ['SKU', 'PN Dell', 'Descripcion', 'Precio USD', 'Stock Ref', 'Pronto Stock']
        rows = [['A', 'PN-A', 'Notebook', 3552, 40, 100],
                headers, ['B', 'PN-B', 'Sin stock', 4800, 0, 200],
                ['C', 'PN-C', 'OPEN BOX', 100, 1, None]]
        summary, df = procesar_archivos([excel(headers, rows)], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(list(df.provider_code), ['A'])
        self.assertEqual(df.iloc[0].quantity, 40)
        self.assertEqual(df.iloc[0].currency_unaware_cost_neto, 3552)
        self.assertEqual(df.iloc[0].mpn, 'PN-A')
        self.assertEqual(df.iloc[0].brand, '')

    def test_intcomex_hp_weekly(self):
        headers = ['Local Sku', 'MPN', 'Product Name (Local)', 'Stock', 'Llegada', 'Precio Normal', 'Precio Promo Octubre', 'Ext Gtia']
        rows = [['A', 'PN-A', 'Notebook', 714, 100, 1409, 1319, 'Precio USD 68'],
                ['B', 'PN-B', 'Tránsito', 0, 600, 100, 90, None],
                ['C', 'PN-C', 'Normal', 1, 0, 105, None, None],
                ['D', 'PN-D', 'BAD BOX monitor', 1, 0, 100, 99, None]]
        summary, df = procesar_archivos([excel(headers, rows)], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(list(df.provider_code), ['A', 'C'])
        self.assertEqual(list(df.currency_unaware_cost_neto), [1319, 105])
        self.assertEqual(list(df.quantity), [714, 1])
        self.assertEqual(df.iloc[0].mpn, 'PN-A')
        self.assertTrue((df.brand == '').all())

    def test_intcomex_logitech_csv(self):
        text = ('Lista de precios,,,,,,,\n'
                'SKU XCL,Logitech Part Number,DESCRIPCIÓN, STOCK , COSTO NETO USD ,PVP SUGERIDO,COSTO NETO USD VOLUMEN,VOLUMEN MINIMO\n'
                'ID010LOG58,910-004053,Mouse M90," 1,019 ",3.6,5390,3.5,90\n'
                'B,PN-B,Sin stock,0,10,20,9,5\n'
                'C,PN-C,BAD BOX,1,10,20,9,5\n'
                'D,PN-D,Comentario,30 DIAS,10,20,9,5\n')
        f = BytesIO(text.encode('utf-8-sig')); f.name = 'Logitech.csv'
        summary, df = procesar_archivos([f], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0].quantity, 1019)
        self.assertEqual(df.iloc[0].currency_unaware_cost_neto, 3.6)
        self.assertEqual(df.iloc[0].mpn, '910-004053')
        self.assertEqual(df.iloc[0].brand, '')

    def test_gtc_volume_and_ean(self):
        headers = ['CODIGO GTC', 'EAN', 'DESCRIPCION', 'STOCK', 'NETO USD', 'OFERTA NETO USD', 'OBS']
        rows = [['A', '123456', 'Producto', 358, 46, 43.5, 'PRECIO X COMPRA DE 150 UNIDADES O MAS'],
                ['B', '234567', 'Otro', 2, 50, 40, None],
                ['C', None, 'Sin normal', 5, None, 20, 'OFERTA POR VOLUMEN']]
        summary, result = procesar_archivos([excel(headers, rows)], FECHA, 'Gtc ribbon')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(list(result.currency_unaware_cost_neto), [46, 40])
        self.assertTrue((result.mpn == '').all())
        self.assertTrue((result.brand == '').all())

    def test_gtc_numeric_pn(self):
        _, result = procesar_archivos([excel(['CODIGO GTC', 'PART NUMBER', 'DESCRIPCION', 'STOCK', 'OFERTA USD NETO'],
            [['107XP00002', 75261462, 'Producto', 48, 32]])], FECHA, 'Gtc ribbon')
        self.assertEqual(result.iloc[0].mpn, '75261462')

    def test_intcomex_top(self):
        headers = ['Marca ', 'Product Sub Category', 'SKU', 'Part Number Marca', 'Product Name', 'Stock', ' Precios Top']
        rows = [['ASUS', 'PSU', ' CS000ASU21 ', '90YE00V2-B0AA00', 'Fuente', 86, 338.26804123711344],
                ['MSI', 'GPU', 'B', 'PN-B', 'Video', 0, 100]]
        summary, result = procesar_archivos([excel(headers, rows)], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0].provider_code, 'CS000ASU21')
        self.assertAlmostEqual(result.iloc[0].currency_unaware_cost_neto, 338.26804123711344)
        self.assertEqual(result.iloc[0].mpn, '90YE00V2-B0AA00')
        self.assertEqual(result.iloc[0].brand, 'ASUS')
        self.assertEqual(result.iloc[0]['name'], 'Fuente')

    def test_tecnoglobal_components_uses_volume_price(self):
        headers = ['CodigoArticulo', 'PartNumber', 'DescripCorta', 'DescripcionMarca', 'Stock ', 'Precio Volumen', 'Precio Venta']
        rows = [['AM0-285', '100-100001721WOF', 'Procesador', 'AMD', 22, 307.3, 312.2],
                ['B', 'PN-B', 'Otro', 'Asus', 0, 10, 12],
                ['C', 'PN-C', 'Sin precio volumen', 'AMD', 1, None, 10]]
        summary, result = procesar_archivos([excel(headers, rows)], FECHA, 'Tecnoglobal')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0].currency_unaware_cost_neto, 307.3)
        self.assertEqual(result.iloc[0].mpn, '100-100001721WOF')
        self.assertEqual(result.iloc[0].brand, 'AMD')

    def test_intcomex_hpe_immediate_only(self):
        headers = ['PN', 'DESCRIPCIÓN', 'PRECIO UN\nVENTA', 'Precio\nPROMO Referencial', 'Tiempo\nENTREGA', 'SKU INTCOMEX']
        rows = [['PN-A', 'Servidor', 3000, 2370, '20 unidades\nENTREGA INMEDIATA', 'A'],
                ['PN-B', 'Tránsito', 3000, 2500, '5 Unidades en transito 7-15 dias', 'B'],
                ['PN-C', 'Fábrica', 3000, 2500, 'pedido fabrica 35-60 dias', 'C'],
                ['PN-D', 'Normal', 1500, None, '1 UN Entrega Inmediata', 'D'],
                ['PN-E', 'Sin cantidad', 1500, 1000, 'Entrega Inmediata', 'E']]
        stream = excel(headers, rows)
        book = openpyxl.load_workbook(stream)
        book.active.title = 'LP GEN 11'
        book.create_sheet('PROMOS E INCENTIVOS', 0)
        stream.seek(0); stream.truncate(0); book.save(stream); stream.seek(0)
        summary, result = procesar_archivos([stream], FECHA, 'Intcomex')
        self.assertEqual(summary.iloc[0]['Estado'], 'Procesado')
        self.assertEqual(list(result.provider_code), ['A', 'D'])
        self.assertEqual(list(result.quantity), [20, 1])
        self.assertEqual(list(result.currency_unaware_cost_neto), [2370, 1500])
        self.assertTrue((result.brand == '').all())


if __name__ == '__main__':
    unittest.main()
